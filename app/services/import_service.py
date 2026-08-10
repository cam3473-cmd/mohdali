"""استيراد بيانات الأعضاء من ملف Excel (سجل الجمعية الحالي) إلى النظام.

يتعامل مع خصوصيات السجلات الورقية/الجدولية الشائعة لدى الجمعيات الأهلية:
تواريخ هجرية أو ميلادية مختلطة (أحيانًا سنة فقط)، وأعمدة قد تختلف تسميتها
قليلًا بين نسخ الملف (مثل "حالة السداد" مقابل "السداد").
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field

import openpyxl
from hijridate import Hijri
from sqlalchemy.orm import Session

from app.db.models import BoardPosition, FeeStatus, Member, MemberStatus, MembershipFee, User
from app.services.audit import log_action
from app.services.bylaw_settings_service import get_settings

# تسميات محتملة لكل عمود (يُطابَق النص بعد إزالة التطويل والمسافات الزائدة)
CANONICAL_HEADERS: dict[str, list[str]] = {
    "full_name": ["الاسم", "الاسم الكامل", "اسم العضو"],
    "membership_number": ["رقم عضوية", "رقم العضوية"],
    "join_date_raw": ["تاريخ بدء العضوية", "تاريخ الانضمام"],
    "member_type": ["نوع العضوية"],
    "position": ["المنصب الحالي", "المنصب"],
    "payment_status": ["حالة السداد", "السداد"],
    "fee_amount": ["قيمة الاشتراك"],
    "active_flag": ["فعال"],
    "receipt_number": ["رقم القيد", "رقم السند"],
    "national_id": ["السجل", "رقم الهوية", "رقم الهوية/السجل"],
    "birth_date_raw": ["تاريخ الميلاد"],
    "phone": ["الجوال", "رقم الجوال"],
    "gender": ["الجنس"],
    "qualification": ["المؤهل"],
    "city": ["المدينة"],
    "occupation": ["العمل"],
}

# مناصب لا تُعتبر منصبًا فعليًا في مجلس الإدارة (تعني عضوًا عاملاً عاديًا فقط)
NON_BOARD_POSITION_VALUES = {"", "عضو عامل", "عضو", "-", "—"}

TRUTHY_ACTIVE_VALUES = {"þ", "true", "1", "نعم", "✓", "yes", "فعال"}
FALSY_ACTIVE_VALUES = {"ý", "false", "0", "لا", "✗", "no", "غير فعال"}

PAID_STATUS_VALUES = {"منتظم", "سدد", "سدد بالوقت", "paid"}


@dataclass
class ImportReport:
    created: int = 0
    updated: int = 0
    skipped: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _normalize_header(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("ـ", "")  # إزالة حرف التطويل
    return " ".join(text.split()).strip()


def list_sheet_names(file_path: str) -> list[str]:
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    return wb.sheetnames


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    max_scan = min(ws.max_row, 10)
    for row_idx in range(1, max_scan + 1):
        col_map: dict[str, int] = {}
        for col_idx, cell in enumerate(ws[row_idx], start=1):
            normalized = _normalize_header(cell.value)
            for field_name, variants in CANONICAL_HEADERS.items():
                if normalized in variants and field_name not in col_map:
                    col_map[field_name] = col_idx
        if "full_name" in col_map:
            return row_idx, col_map
    raise ValueError("تعذر العثور على صف العناوين في الملف (لم يتم التعرف على عمود الاسم)")


def _parse_year_only(value) -> dt.date | None:
    """يحوّل قيمة سنة فقط (هجرية إن كانت صغيرة، وإلا ميلادية) إلى تاريخ تقريبي (1/1)."""
    try:
        year = int(float(value))
    except (TypeError, ValueError):
        return None
    if year <= 0:
        return None
    if year <= 1500:  # سنة هجرية غالبًا
        try:
            g = Hijri(year, 1, 1).to_gregorian()
            return dt.date(g.year, g.month, g.day)
        except Exception:
            return None
    return dt.date(year, 1, 1)


def _parse_flexible_date(value) -> dt.date | None:
    """يحوّل تاريخًا قد يكون: كائن تاريخ ميلادي جاهز، نص هجري (فواصل /)، أو سنة فقط."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        return _parse_year_only(value)

    text = str(value).strip()
    if not text:
        return None
    if "/" in text or "-" in text:
        sep = "/" if "/" in text else "-"
        parts = [p for p in text.split(sep) if p.strip()]
        try:
            nums = [int(p) for p in parts]
        except ValueError:
            return None
        if len(nums) != 3:
            return None
        # حدد الرقم الذي يمثل السنة (القيمة الكبيرة > 31، لأن اليوم/الشهر لا يتجاوزان ذلك)
        year_idx = max(range(3), key=lambda i: nums[i])
        year = nums[year_idx]
        remaining = [n for i, n in enumerate(nums) if i != year_idx]
        month, day = (remaining[0], remaining[1]) if year_idx == 0 else (remaining[1], remaining[0])
        month = max(1, min(month, 12))
        day = max(1, min(day, 28))
        try:
            g = Hijri(year, month, day).to_gregorian()
            return dt.date(g.year, g.month, g.day)
        except Exception:
            return None
    return _parse_year_only(text)


def _interpret_active_flag(value) -> bool:
    text = _normalize_header(value).lower()
    if text in FALSY_ACTIVE_VALUES:
        return False
    if text in TRUTHY_ACTIVE_VALUES:
        return True
    return True  # افتراضيًا نعتبر العضو فعالًا ما لم يُذكر صراحة أنه غير ذلك


def import_members_from_excel(
    session: Session,
    actor: User,
    file_path: str,
    sheet_name: str | None = None,
    fee_year: int | None = None,
) -> ImportReport:
    report = ImportReport()
    fee_year = fee_year or dt.date.today().year

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    header_row, col_map = _find_header_row(ws)

    settings = get_settings(session)
    imported_member_types: set[str] = set()

    for row_idx in range(header_row + 1, ws.max_row + 1):
        row_cells = ws[row_idx]

        def get(field_name: str):
            col = col_map.get(field_name)
            return row_cells[col - 1].value if col else None

        full_name = _normalize_header(get("full_name"))
        if not full_name:
            continue

        membership_number = _normalize_header(get("membership_number")) or None
        national_id = _normalize_header(get("national_id")) or None

        existing = None
        if national_id:
            existing = session.query(Member).filter(Member.national_id_or_cr == national_id).first()
        if existing is None and membership_number:
            existing = session.query(Member).filter(Member.membership_number == membership_number).first()

        member_type = _normalize_header(get("member_type")) or "عادية"
        imported_member_types.add(member_type)

        join_date = _parse_flexible_date(get("join_date_raw")) or dt.date.today()
        birth_date = _parse_flexible_date(get("birth_date_raw"))
        phone_raw = get("phone")
        phone = str(int(phone_raw)) if isinstance(phone_raw, float) else (str(phone_raw).strip() if phone_raw else None)
        gender = _normalize_header(get("gender")) or None
        qualification = _normalize_header(get("qualification")) or None
        city = _normalize_header(get("city")) or None
        occupation = _normalize_header(get("occupation")) or None
        is_active = _interpret_active_flag(get("active_flag"))

        if existing is None:
            member = Member(
                membership_number=membership_number,
                national_id_or_cr=national_id,
                full_name=full_name,
                member_type=member_type,
                join_date=join_date,
                birth_date=birth_date,
                phone=phone,
                gender=gender,
                qualification=qualification,
                city=city,
                occupation=occupation,
                status=MemberStatus.ACTIVE if is_active else MemberStatus.SUSPENDED,
            )
            session.add(member)
            session.flush()
            log_action(session, actor, "import_member_created", "member", member.id, details=f"row={row_idx}")
            report.created += 1
        else:
            member = existing
            member.full_name = full_name
            member.member_type = member_type
            member.membership_number = membership_number or member.membership_number
            member.national_id_or_cr = national_id or member.national_id_or_cr
            member.join_date = join_date
            member.birth_date = birth_date or member.birth_date
            member.phone = phone or member.phone
            member.gender = gender or member.gender
            member.qualification = qualification or member.qualification
            member.city = city or member.city
            member.occupation = occupation or member.occupation
            member.status = MemberStatus.ACTIVE if is_active else MemberStatus.SUSPENDED
            log_action(session, actor, "import_member_updated", "member", member.id, details=f"row={row_idx}")
            report.updated += 1

        payment_status = _normalize_header(get("payment_status"))
        fee_amount = get("fee_amount") or 0
        if payment_status in PAID_STATUS_VALUES and fee_amount:
            fee = (
                session.query(MembershipFee)
                .filter(MembershipFee.member_id == member.id, MembershipFee.fee_year == fee_year)
                .first()
            )
            if fee is None:
                fee = MembershipFee(member_id=member.id, fee_year=fee_year)
                session.add(fee)
            fee.amount = float(fee_amount)
            fee.status = FeeStatus.PAID
            receipt = get("receipt_number")
            fee.receipt_number = str(int(receipt)) if isinstance(receipt, float) else (str(receipt) if receipt else None)

        position_title = _normalize_header(get("position"))
        if position_title and position_title not in NON_BOARD_POSITION_VALUES:
            has_current = (
                session.query(BoardPosition)
                .filter(
                    BoardPosition.member_id == member.id,
                    BoardPosition.title == position_title,
                    BoardPosition.end_date.is_(None),
                )
                .first()
            )
            if not has_current:
                session.add(BoardPosition(member_id=member.id, title=position_title, start_date=None))

        if not birth_date and get("birth_date_raw"):
            report.warnings.append(f"سطر {row_idx} ({full_name}): تعذر تفسير تاريخ الميلاد '{get('birth_date_raw')}'")

    session.commit()

    missing_voting_types = imported_member_types - set(settings.voting_member_types)
    if missing_voting_types:
        report.warnings.append(
            "أنواع عضوية مستوردة غير مدرجة ضمن 'أنواع العضوية المصوّتة' الحالية في الإعدادات: "
            + "، ".join(sorted(missing_voting_types))
            + " — راجع شاشة الإعدادات إذا كان يجب أن تمنح هذه الأنواع حق التصويت."
        )

    return report


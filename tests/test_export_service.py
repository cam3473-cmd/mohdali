from datetime import date

import openpyxl

from app.services import board_service, export_service, import_service, membership_service


def test_export_members_to_excel_writes_expected_data(db_session, admin_user, tmp_path):
    member = membership_service.submit_membership_request(
        db_session,
        admin_user,
        full_name="عضو للتصدير",
        member_type="عادية",
        membership_number="9",
        national_id_or_cr="1046952857",
        join_date=date(2020, 1, 1),
        birth_date=date(1980, 5, 1),
        phone="500000000",
        gender="ذكر",
        qualification="جامعي",
        city="السليل",
        occupation="موظف",
    )
    membership_service.approve_membership(db_session, admin_user, member)
    membership_service.record_fee_payment(db_session, admin_user, member, fee_year=date.today().year, amount=300)
    board_service.assign_position(db_session, admin_user, member, title="أمين الصندوق")

    output_path = str(tmp_path / "export.xlsx")
    export_service.export_members_to_excel(db_session, [member], output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == export_service.EXPORT_HEADERS
    row = [cell.value for cell in ws[2]]
    data = dict(zip(header, row))
    assert data["الاسم"] == "عضو للتصدير"
    assert data["السجل"] == "1046952857"
    assert data["المنصب الحالي"] == "أمين الصندوق"
    assert data["السداد"] == "منتظم"
    assert data["فعال"] == "نعم"


def test_export_then_reimport_round_trip(db_session, admin_user, tmp_path):
    member = membership_service.submit_membership_request(
        db_session,
        admin_user,
        full_name="عضو الجولة الكاملة",
        member_type="عادية",
        membership_number="10",
        national_id_or_cr="1099999999",
        join_date=date(2019, 3, 15),
    )
    membership_service.approve_membership(db_session, admin_user, member)

    output_path = str(tmp_path / "roundtrip.xlsx")
    export_service.export_members_to_excel(db_session, [member], output_path)

    report = import_service.import_members_from_excel(db_session, admin_user, output_path)
    assert report.created == 0
    assert report.updated == 1
